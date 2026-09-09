from docx import Document
from docx.shared import Pt
from pathlib import Path

md_path = Path('DATABASE_SCHEMA.md')
if not md_path.exists():
    print('DATABASE_SCHEMA.md not found in current folder')
    raise SystemExit(1)

text = md_path.read_text(encoding='utf-8')
lines = text.splitlines()

doc = Document()
code_block = False
code_lines = []

for line in lines:
    if line.strip().startswith('```'):
        code_block = not code_block
        if not code_block and code_lines:
            p = doc.add_paragraph()
            run = p.add_run('\n'.join(code_lines))
            run.font.name = 'Courier New'
            run.font.size = Pt(9)
            code_lines = []
        continue

    if code_block:
        code_lines.append(line)
        continue

    stripped = line.lstrip()
    if not stripped:
        doc.add_paragraph('')
        continue

    # Headings
    if stripped.startswith('# '):
        doc.add_heading(stripped[2:].strip(), level=1)
        continue
    if stripped.startswith('## '):
        doc.add_heading(stripped[3:].strip(), level=2)
        continue
    if stripped.startswith('### '):
        doc.add_heading(stripped[4:].strip(), level=3)
        continue

    # Horizontal rule
    if stripped.startswith('---'):
        p = doc.add_paragraph('\u2014' * 10)
        continue

    # List items
    if stripped.startswith('- '):
        doc.add_paragraph(stripped[2:].strip(), style='List Bullet')
        continue

    # Code fence or mermaid block fallback
    if stripped.startswith('```'):
        continue

    # Default paragraph
    doc.add_paragraph(stripped)

# Save
out = Path('DATABASE_SCHEMA.docx')
doc.save(out)
print(f'Wrote {out.resolve()}')
